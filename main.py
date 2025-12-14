# hi
# PLEASE USE "pip install python-docx" IF THERE IS AN ERROR WITH THE DOCX IMPORT
# IMPORTS - Valerie
from docx import Document
from docx.shared import Pt
document = Document()

# OPTIONS AND INPUTS - Alvin
name = input("Enter your Name: ")
job_title = input("Enter your Job Title: ")
date = input("Enter today's Date (MM/DD/YYYY): ")
room_number = input("Enter your Room Number: ")

department = input("Select Department (1 = HR, 2 = IT, 3 = Finance): ")

if department == "1":
    department_name = "HR"
elif department == "2":
    department_name = "IT"
elif department == "3":
    department_name = "Finance"
else:
    input("Invalid option. Please enter 1, 2, or 3: ")

print(department_name)

# WORD COMPONENTS - Arshia
email_text = (
    f"Good day {department_name} Department,\n"
    f"This is {name}, {job_title}, requesting assistance.\n"
    f"Please email back and refer to room {room_number} as soon as possible.\n"
    f"This ticket was opened on {date}.\n"
    f"Thank you,\n"
    f"{name}\n"
    f"Room {room_number}"
)
email_heading = (
    f"Email Ticket for {department_name}- {date}"
)

print("Loading preview of your email to be put in a word doc...")
if department_name == "HR":
    document.add_heading(email_heading, level=1)
    print("------YOUR EMAIL PREVIEW------")
    print(email_text)
    print("------------------------------")
    document.add_paragraph(email_text)
elif department_name == "IT":
    document.add_heading(email_heading, level=1)
    print("------YOUR EMAIL PREVIEW------")
    print(email_text)
    print("------------------------------")
    document.add_paragraph(email_text)
elif department_name == "Finance":
    document.add_heading(email_heading, level=1)
    print("------YOUR EMAIL PREVIEW------")
    print(email_text)
    print("------------------------------")
    document.add_paragraph(email_text)
else:
    print("Invalid. Please restart the program.")

# SAVING AND NAMING - Valerie
document.save(f"{name}_ticket.docx")